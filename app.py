# app.py

# Requirements:
# pip install -U gradio google-genai PyPDF2 python-docx pillow pandas openpyxl
# Optional for OCR/scanned PDFs:
# pip install pdf2image pytesseract
# Ensure Tesseract binary is installed and on PATH for pytesseract.

import os
import re
import json
import time
import random
import logging
from typing import Dict, List, Any, Tuple, Optional
from pathlib import Path

import gradio as gr

# Document processing
import PyPDF2
from docx import Document
from PIL import Image
import pandas as pd  # noqa: F401
from openpyxl import load_workbook  # noqa: F401

# Optional OCR/scanned-pdf support
try:
    from pdf2image import convert_from_path
    PDF2IMAGE_AVAILABLE = True
except Exception:
    PDF2IMAGE_AVAILABLE = False

try:
    import pytesseract
    OCR_AVAILABLE = True
except Exception:
    OCR_AVAILABLE = False

# Google Gen AI SDK (current)
GENAI_CLIENT: Optional[Any] = None
GENAI_MODEL_NAME: Optional[str] = None
try:
    from google import genai
    GENAI_LIB_AVAILABLE = True
except Exception:
    GENAI_LIB_AVAILABLE = False

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("advanced_legal_analyzer")

APP_TITLE = "⚖️ Advanced AI Legal Analyzer"
APP_TAGLINE = "Instant, intelligent insights into legal documents."

# -----------------------------
# Gemini client + model resolver
# -----------------------------
def init_gemini_client_and_model() -> Tuple[Optional[Any], Optional[str]]:
    """
    Initialize the Google Gen AI client and determine a valid model that supports generateContent.
    Prefer stable 1.5 families and avoid experimental/learnlm models.
    Returns (client, model_name) or (None, None) if unavailable.
    """
    if not GENAI_LIB_AVAILABLE:
        logger.warning("google-genai library not available. AI features will be disabled.")
        return None, None

    api_key = os.environ.get("GEMINI_API_KEY") or os.environ.get("GOOGLE_API_KEY")
    if not api_key:
        logger.warning("GEMINI_API_KEY is not set. AI features will be disabled.")
        return None, None

    try:
        client = genai.Client(api_key=api_key)

        def list_generate_capable(skip_experimental: bool) -> List[str]:
            names: List[str] = []
            for m in client.models.list():
                actions = set(getattr(m, "supported_actions", []) or [])
                name = getattr(m, "name", "") or ""
                if "generateContent" in actions and name:
                    lname = name.lower()
                    if skip_experimental and any(tag in lname for tag in ["experimental", "learnlm", "2.0-flash-experimental"]):
                        continue
                    names.append(name)
            return names

        # First try stable
        generate_capable = list_generate_capable(skip_experimental=True)
        if not generate_capable:
            # Fallback: allow any generate-capable
            generate_capable = list_generate_capable(skip_experimental=False)

        if not generate_capable:
            logger.error("No models supporting generateContent available for this key/project.")
            return client, None

        def pick(preferences: List[str]) -> Optional[str]:
            for pref in preferences:
                cands = [n for n in generate_capable if pref in n]
                if cands:
                    return sorted(cands)[-1]
            return None

        # Prefer stable families
        model_name = pick(["gemini-1.5-pro", "gemini-1.5-flash"]) or sorted(generate_capable)[-1]
        logger.info(f"Selected Gemini model: {model_name}")
        return client, model_name
    except Exception as e:
        logger.error(f"Failed to initialize Gemini client: {e}")
        return None, None

# -----------------------------
# Analyzer
# -----------------------------
class AdvancedLegalAnalyzer:
    def __init__(self, client: Optional[Any], model_name: Optional[str]):
        self.client = client
        self.model_name = model_name

        self.risk_keywords = {
            'critical_risk': {
                'terms': [
                    'irrevocable', 'unlimited liability', 'perpetual obligation',
                    'waive all rights', 'complete indemnification', 'liquidated damages',
                    'penalty clause', 'criminal liability', 'personal guarantee',
                    'joint and several liability', 'unlimited damages', 'unilateral termination'
                ],
                'weight': 15, 'description': 'Critical terms that could expose you to significant legal or financial risk'
            },
            'high_risk': {
                'terms': [
                    'indemnify', 'hold harmless', 'at will termination', 'sole discretion',
                    'no warranty', 'force majeure', 'assignment of rights',
                    'non-compete', 'confidentiality breach', 'automatic renewal', 'binding arbitration'
                ],
                'weight': 10, 'description': 'High-risk terms that significantly limit rights or increase obligations'
            },
            'medium_risk': {
                'terms': [
                    'termination', 'breach', 'default', 'damages', 'liability',
                    'jurisdiction', 'governing law', 'modification', 'amendment', 'notice period'
                ],
                'weight': 6, 'description': 'Medium-risk terms that require careful consideration'
            },
            'favorable_terms': {
                'terms': [
                    'mutual agreement', 'reasonable notice', 'good faith negotiation',
                    'proportional liability', 'cap on damages', 'right to cure',
                    'mutual termination', 'equitable relief', 'cost sharing'
                ],
                'weight': -3, 'description': 'Generally favorable terms that protect interests'
            }
        }

        self.document_patterns = {
            'government_form': ['आय', 'घोषणा पत्र', 'सरकार', 'प्रमाण पत्र', 'form', 'affidavit', 'declaration'],
            'employment_contract': ['employee', 'employer', 'salary', 'benefits', 'job description', 'termination', 'non-compete', 'confidentiality'],
            'rental_agreement': ['landlord', 'tenant', 'lease', 'rent', 'security deposit', 'premises', 'utilities', 'maintenance', 'eviction'],
            'service_agreement': ['service provider', 'client', 'deliverables', 'scope of work', 'payment terms', 'intellectual property'],
        }

    # --------
    # Extractors
    # --------
    def extract_text_from_file(self, file_path: str) -> Tuple[str, Dict]:
        file_path_obj = Path(file_path)
        metadata = {
            'file_name': file_path_obj.name,
            'file_size': file_path_obj.stat().st_size if file_path_obj.exists() else 0,
            'file_type': file_path_obj.suffix.lower()
        }

        text = ""
        ext = metadata['file_type']

        try:
            if ext == '.pdf':
                text, pdf_meta = self._extract_from_pdf_hybrid(file_path)
                metadata.update(pdf_meta)
            elif ext == '.docx':
                doc = Document(file_path)
                text = "\n".join([para.text for para in doc.paragraphs])
            elif ext in ['.png', '.jpg', '.jpeg']:
                if not OCR_AVAILABLE:
                    raise ValueError("OCR is not available. Install pytesseract for image analysis.")
                text, img_meta = self._extract_from_image(file_path_obj)
                metadata.update(img_meta)
            else:
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    text = f.read()
        except Exception as e:
            logger.error(f"Extraction error: {e}")
            raise

        text = self._advanced_text_cleaning(text)
        metadata['word_count'] = len(text.split())
        return text, metadata

    def _extract_from_pdf_hybrid(self, file_path: str) -> Tuple[str, Dict]:
        text = ""
        metadata = {'pages': 0, 'ocr_pages': []}
        try:
            reader = PyPDF2.PdfReader(file_path)
            metadata['pages'] = len(reader.pages)
            for idx, page in enumerate(reader.pages):
                page_text = ""
                try:
                    page_text = page.extract_text() or ""
                except Exception:
                    page_text = ""

                if len(page_text.strip()) < 50 and PDF2IMAGE_AVAILABLE and OCR_AVAILABLE:
                    try:
                        images = convert_from_path(file_path, first_page=idx + 1, last_page=idx + 1)
                        if images:
                            page_text = pytesseract.image_to_string(images[0], lang='eng+hin')
                            metadata['ocr_pages'].append(idx + 1)
                    except Exception as ocr_err:
                        logger.warning(f"OCR fallback failed on page {idx+1}: {ocr_err}")

                text += f"\n--- Page {idx + 1} ---\n{page_text}\n"
        except Exception as e:
            logger.error(f"PDF extraction error: {e}")
            raise
        return text, metadata

    def _extract_from_image(self, file_path: Path) -> Tuple[str, Dict]:
        try:
            with Image.open(file_path) as img:
                t = pytesseract.image_to_string(img, lang='eng+hin')
                return t, {'image_resolution': f"{img.width}x{img.height}"}
        except Exception as e:
            logger.error(f"Image OCR error: {e}")
            raise

    @staticmethod
    def _advanced_text_cleaning(text: str) -> str:
        text = re.sub(r'\s+', ' ', text)
        text = re.sub(r'--- Page \d+ ---', '', text)
        return text.strip()

    # -------------
    # Heuristics
    # -------------
    def identify_document_type(self, text: str) -> Tuple[str, float]:
        if not text:
            return "Unknown", 0.0
        scores: Dict[str, int] = {}
        for doc_type, keywords in self.document_patterns.items():
            scores[doc_type] = sum(text.lower().count(k) for k in keywords)
        total = sum(scores.values())
        if total == 0:
            return "General Document", 0.5
        best = max(scores, key=scores.get)
        return best.replace('_', ' ').title(), scores[best] / total

    def _enhanced_risk_assessment(self, text: str) -> Dict[str, Any]:
        findings: Dict[str, Any] = {}
        total = 0
        tl = text.lower()
        for level, data in self.risk_keywords.items():
            found = [{'term': term, 'count': tl.count(term)} for term in data['terms'] if term in tl]
            if found:
                findings[level] = {'terms_found': found, 'description': data['description']}
                total += sum(t['count'] * data['weight'] for t in found)

        if total >= 100:
            overall = "Critical"
        elif total >= 50:
            overall = "High"
        elif total >= 20:
            overall = "Medium"
        elif total > -10:
            overall = "Low"
        else:
            overall = "Favorable"
        return {'overall_risk_level': overall, 'risk_score': total, 'risk_findings': findings}

    def _extract_key_terms(self, text: str) -> List[Dict]:
        legal_terms_db = {
            'indemnification': 'One party agrees to cover the losses of another.',
            'limitation of liability': 'Caps the amount of damages a party can be responsible for.',
            'confidentiality': 'Obligation to keep certain information secret.',
            'termination': 'Conditions under which the agreement can be ended.',
            'governing law': 'Specifies which laws will be used to interpret the agreement.',
            'arbitration': 'A private method of resolving disputes outside of court.'
        }
        tl = text.lower()
        return [{'term': term.title(), 'definition': definition}
                for term, definition in legal_terms_db.items() if term in tl]

    def _enhanced_compliance_check(self, text: str) -> Dict[str, Any]:
        clauses = {
            "Termination": ['terminate', 'termination'],
            "Governing Law": ['governed by'],
            "Dispute Resolution": ['arbitration', 'mediation'],
            "Confidentiality": ['confidential'],
            "Liability": ['liability', 'indemnify']
        }
        tl = text.lower()
        results = {cl: any(kw in tl for kw in kws) for cl, kws in clauses.items()}
        score = int((sum(results.values()) / len(clauses)) * 100)
        return {'score': score, 'checklist': results}

    def _calculate_overall_score(self, analysis: Dict) -> Dict[str, Any]:
        risk_score = analysis.get('risk_score', 0)
        compliance_score = analysis.get('compliance', {}).get('score', 0)
        normalized_risk_score = max(0, 100 - (risk_score * 0.66))
        final_score = int((normalized_risk_score * 0.6) + (compliance_score * 0.4))
        if final_score >= 90:
            grade = "A+"
        elif final_score >= 80:
            grade = "A"
        elif final_score >= 70:
            grade = "B"
        elif final_score >= 60:
            grade = "C"
        elif final_score >= 50:
            grade = "D"
        else:
            grade = "F"
        return {'score': final_score, 'grade': grade}

    # -------------
    # AI assistance with retry/backoff
    # -------------
    def _call_gemini(self, prompt: str) -> str:
        if not self.client or not self.model_name:
            raise RuntimeError("AI model not available. Configure GEMINI_API_KEY and restart.")
        delays = [1.5, 3, 6, 10]  # exponential-ish backoff
        last_err: Optional[Exception] = None
        for d in delays:
            try:
                resp = self.client.models.generate_content(model=self.model_name, contents=prompt)
                return getattr(resp, "text", "") or ""
            except Exception as e:
                msg = str(e).lower()
                if any(tok in msg for tok in ["503", "unavailable", "overloaded", "429", "resource_exhausted"]):
                    time.sleep(d + random.uniform(0, 0.5))
                    last_err = e
                    continue
                raise
        raise RuntimeError(f"Gemini request failed after retries: {last_err}")

    def _enhanced_ai_analysis(self, text: str, doc_type: str) -> Dict[str, Any]:
        try:
            prompt = f"""
As an expert analyst reviewing a '{doc_type}', provide a detailed JSON response.
Analyze the document text. Be concise.

Respond with this exact JSON structure:
{{
  "document_title": "A concise, likely title for this document.",
  "parties_involved": "A brief description of parties. If a contract, name Party A/B roles. If a form, name Applicant/Authority.",
  "executive_summary": "A 3-4 sentence summary for a busy executive, highlighting the purpose and key outcomes.",
  "eli5_summary": "Explain Like I'm 5: A one-sentence, super-simple explanation.",
  "key_financial_terms": ["List key monetary values found, e.g., 'Total Income: ₹2,00,000'"],
  "important_dates": ["List key dates found, e.g., 'Date of Issue: 18/08/25'"],
  "potential_issues": [
    {{"issue": "Describe a potential problem", "severity": "High|Medium|Low", "recommendation": "Suggest an action"}}
  ],
  "action_items": ["List immediate, actionable steps for the user."]
}}

Document text: "{text[:18000]}"
"""
            response_text = self._call_gemini(prompt).strip()
            json_match = re.search(r"``````", response_text, re.DOTALL)
            if not json_match:
                json_match = re.search(r"(\{.*\})", response_text, re.DOTALL)
            if not json_match:
                raise ValueError("No valid JSON object found in AI response.")
            json_text = json_match.group(1)
            return json.loads(json_text)
        except Exception as e:
            logger.warning(f"AI analysis failed: {e}")
            return {"executive_summary": f"AI analysis failed: {e}"}

    # -------------
    # Orchestrator
    # -------------
    def analyze_document_comprehensive(self, text: str, metadata: Dict) -> Dict[str, Any]:
        doc_type, confidence = self.identify_document_type(text)
        risk = self._enhanced_risk_assessment(text)
        ai_block = self._enhanced_ai_analysis(text, doc_type) if (self.client and self.model_name) else {}
        analysis = {
            **risk,
            **ai_block,
            'document_type': doc_type,
            'type_confidence': confidence,
            'key_terms': self._extract_key_terms(text),
            'compliance': self._enhanced_compliance_check(text),
            'metadata': metadata
        }
        analysis['overall_score'] = self._calculate_overall_score(analysis)
        return analysis

# -----------------------------
# UI helpers
# -----------------------------
def format_dashboard(analysis: Dict) -> str:
    if not analysis or 'overall_score' not in analysis:
        return "<h3>Analysis Failed</h3>"
    scorecard = analysis.get('overall_score', {})
    risk_level = analysis.get('overall_risk_level', 'N/A')
    risk_colors = {
        "Critical": "#e53e3e", "High": "#dd6b20", "Medium": "#d69e2e",
        "Low": "#38a169", "Favorable": "#3182ce"
    }
    risk_color = risk_colors.get(risk_level, "#718096")
    financial_html = "".join([f'<li>{item}</li>' for item in (analysis.get('key_financial_terms') or [])])
    dates_html = "".join([f'<li>{item}</li>' for item in (analysis.get('important_dates') or [])])
    return f"""
    <div style="padding: 20px; border-radius: 10px; background-color: #f7fafc;">
      <h2 style="margin-top:0;">{analysis.get('document_title', 'Document Analysis')}</h2>
      <p><em>{analysis.get('eli5_summary', '')}</em></p>
      <div style="display: grid; grid-template-columns: repeat(3, 1fr); gap: 15px; text-align: center; margin: 25px 0;">
        <div><h3 style="margin:0;color:#4a5568;">Grade</h3><p style="font-size:2.5em;margin:5px 0;color:{risk_color};font-weight:bold;">{scorecard.get('grade','N/A')}</p></div>
        <div><h3 style="margin:0;color:#4a5568;">Risk Level</h3><p style="font-size:2.5em;margin:5px 0;color:{risk_color};font-weight:bold;">{risk_level}</p></div>
        <div><h3 style="margin:0;color:#4a5568;">Compliance</h3><p style="font-size:2.5em;margin:5px 0;color:#4a5568;font-weight:bold;">{analysis.get('compliance',{}).get('score',0)}%</p></div>
      </div>
      <h3>Executive Summary</h3><p>{analysis.get('executive_summary','Not available.')}</p>
      <h3>Parties Involved</h3><p>{analysis.get('parties_involved','Not identified.')}</p>
      <div style="display:grid; grid-template-columns: 1fr 1fr; gap: 20px;">
        <div><h3>💰 Key Financial Terms</h3><ul>{financial_html or '<li>None identified</li>'}</ul></div>
        <div><h3>🗓️ Important Dates</h3><ul>{dates_html or '<li>None identified</li>'}</ul></div>
      </div>
      <h3>Action Items</h3><ul>{''.join([f'<li>{item}</li>' for item in (analysis.get('action_items') or [])]) or '<li>None listed</li>'}</ul>
    </div>
    """

def format_risk_analysis(analysis: Dict) -> Tuple[str, str]:
    if not analysis or 'risk_findings' not in analysis:
        return "<p>No risk data.</p>", "<p>No risk data.</p>"
    rf = analysis.get('risk_findings', {})
    risk_data = {
        "Critical": len(rf.get('critical_risk', {}).get('terms_found', [])),
        "High": len(rf.get('high_risk', {}).get('terms_found', [])),
        "Medium": len(rf.get('medium_risk', {}).get('terms_found', []))
    }
    risk_colors = {"Critical": "#e53e3e", "High": "#dd6b20", "Medium": "#d69e2e"}
    counts: List[int] = [int(x) for x in risk_data.values()]
    max_val = max(counts) if any(counts) else 1
    chart_html = '<div style="padding:10px;">' + "".join([
        f'<div style="margin-bottom:8px;"><strong style="display:inline-block;width:80px;">{L}</strong>'
        f'<div style="display:inline-block;width:calc(100% - 110px);background:#e2e8f0;">'
        f'<div style="width:{(c/max_val*100)}%;background:{risk_colors[L]};color:white;'
        f'text-align:right;padding:4px;">{c}</div></div></div>'
        for L, c in risk_data.items()
    ]) + '</div>'

    details_md = "## Risk Details\n"
    for L, D in rf.items():
        details_md += f"### {L.replace('_',' ').title()}\n"
        details_md += f"*{D['description']}*\n"
        for t in (D.get('terms_found') or []):
            details_md += f"- **{t['term']}** (found {t['count']} time(s))\n"
    return chart_html, details_md

def format_compliance_analysis(analysis: Dict) -> str:
    if not analysis or 'compliance' not in analysis:
        return "No data."
    checklist = (analysis.get('compliance') or {}).get('checklist') or {}
    return "## Standard Clause Checklist\n" + "\n".join(
        [f"- {'✅' if present else '❌'} **{clause}:** {'Present' if present else 'Missing'}"
         for clause, present in checklist.items()]
    )

def format_key_terms(analysis: Dict) -> str:
    if not analysis or not analysis.get('key_terms'):
        return "No key terms identified."
    return "## Important Legal Terms\n" + "\n".join(
        [f"### {t['term']}\n> {t['definition']}\n" for t in (analysis.get('key_terms') or [])]
    )

# -----------------------------
# Q&A helper
# -----------------------------
def ask_enhanced_question(client, model_name, doc_text: str, question: str, chat_history: List, language: str, analysis_data: Dict):
    chat_history = chat_history or []
    if not doc_text or not question:
        chat_history.append({"role": "user", "content": question})
        chat_history.append({"role": "assistant", "content": "Please upload and analyze a document first."})
        return chat_history

    lang_instruction = "Provide a helpful, conversational answer in simple English."
    if language == "Hinglish":
        lang_instruction = "Provide a helpful, conversational answer in Hinglish (mix of Hindi and English using Roman script)."

    try:
        if not client or not model_name:
            raise RuntimeError("AI model is not available for Q&A.")
        prompt = f"""
You are a helpful assistant explaining a document.

CONTEXT FROM INITIAL ANALYSIS:
- Document Type: {analysis_data.get('document_type', 'N/A')}
- Overall Risk: {analysis_data.get('overall_risk_level', 'N/A')}
- Summary: {analysis_data.get('executive_summary', 'N/A')}

FULL DOCUMENT TEXT (for reference):
---
{doc_text[:15000]}
---

USER QUESTION: "{question}"

INSTRUCTIONS:
- Answer the user's question based on the document text.
- {lang_instruction}
- If the answer is not present, say clearly that it's not found; do not invent.
"""
        # Retry wrapper from analyzer to handle 503/429 is not available here directly;
        # keep it simple and rely on single attempt for chat to avoid long UI waits.
        resp = client.models.generate_content(model=model_name, contents=prompt)
        answer = (getattr(resp, "text", "") or "").strip() or "Sorry, I could not generate an answer."
    except Exception as e:
        logger.error(f"Q&A Error: {e}")
        answer = f"Sorry, an error occurred while answering: {e}"

    chat_history.append({"role": "user", "content": question})
    chat_history.append({"role": "assistant", "content": answer})
    return chat_history

# -----------------------------
# Gradio app
# -----------------------------
def create_enhanced_interface():
    global GENAI_CLIENT, GENAI_MODEL_NAME
    if GENAI_CLIENT is None and GENAI_MODEL_NAME is None:
        client, model_name = init_gemini_client_and_model()
        GENAI_CLIENT, GENAI_MODEL_NAME = client, model_name
        if client and model_name:
            logger.info("Google Gen AI initialized successfully.")
        else:
            logger.warning("AI features will be limited (no model available).")

    with gr.Blocks(css="footer {display: none !important}") as interface:
        gr.HTML(f"""<div style="background:linear-gradient(135deg,#667eea 0%,#764ba2 100%);padding:30px;text-align:center;color:white;">
            <h1 style="margin:0;font-size:2.2em;">{APP_TITLE}</h1>
            <p style="font-size:1.05em;">{APP_TAGLINE}</p></div>""")

        with gr.Row():
            with gr.Column(scale=1, min_width=350):
                file_input = gr.File(label="Upload Document", file_types=[".pdf", ".docx", ".txt", ".png", ".jpg", ".jpeg"])
                analyze_btn = gr.Button("🚀 Analyze Document", variant="primary")
                with gr.Accordion("Document Metadata", open=False):
                    metadata_output = gr.JSON(label="File Info")
                gr.Markdown(
                    f"AI status: {'Ready' if (GENAI_CLIENT and GENAI_MODEL_NAME) else 'Disabled'}"
                    + (f" (model: {GENAI_MODEL_NAME or 'n/a'})" if GENAI_MODEL_NAME else "")
                )

            with gr.Column(scale=3):
                with gr.Tabs() as result_tabs:
                    with gr.TabItem("📊 Dashboard", id=0):
                        dashboard_output = gr.HTML("### Upload a document to see the analysis dashboard.")
                    with gr.TabItem("⚠️ Risk Analysis", id=1):
                        risk_chart_output = gr.HTML()
                        ai_issues_output = gr.JSON(label="AI Identified Issues")
                        risk_details_output = gr.Markdown()
                    with gr.TabItem("✅ Compliance", id=2):
                        compliance_output = gr.Markdown()
                    with gr.TabItem("📚 Key Terms", id=3):
                        key_terms_output = gr.Markdown()
                    with gr.TabItem("💬 Ask a Question", id=4):
                        chatbot = gr.Chatbot(label="Chat about this Document", height=400, type="messages")
                        with gr.Row():
                            lang_select = gr.Radio(["English", "Hinglish"], label="Response Language", value="English", scale=1)
                            question_input = gr.Textbox(
                                label="Your Question",
                                placeholder="e.g., What is the total income mentioned?",
                                scale=3,
                                container=False
                            )
                        ask_btn = gr.Button("Submit Question")

        analysis_state = gr.State({})
        doc_text_state = gr.State("")

        def run_analysis(file, progress=gr.Progress(track_tqdm=True)):
            if file is None:
                raise gr.Error("Please upload a file first.")
            if not PDF2IMAGE_AVAILABLE or not OCR_AVAILABLE:
                gr.Info("Scanned PDF/image OCR may be limited (install pdf2image + pytesseract for full capability).")
            progress(0.1, desc="Initializing...")
            analyzer = AdvancedLegalAnalyzer(GENAI_CLIENT, GENAI_MODEL_NAME)

            progress(0.3, desc="Reading and processing document...")
            text, metadata = analyzer.extract_text_from_file(file.name)
            if not text or len(text.strip()) < 50:
                raise gr.Error("Failed to extract sufficient text. The document may be empty, corrupted, or fully image-based without OCR tools.")

            progress(0.6, desc="Running analysis...")
            analysis = analyzer.analyze_document_comprehensive(text, metadata)

            progress(0.9, desc="Formatting results...")
            dashboard = format_dashboard(analysis)
            chart, risk_details = format_risk_analysis(analysis)
            compliance = format_compliance_analysis(analysis)
            key_terms = format_key_terms(analysis)
            ai_issues = analysis.get('potential_issues', [])

            return (
                dashboard, chart, risk_details, ai_issues,
                compliance, key_terms, metadata,
                analysis, text
            )

        analyze_btn.click(
            fn=run_analysis,
            inputs=[file_input],
            outputs=[
                dashboard_output, risk_chart_output, risk_details_output, ai_issues_output,
                compliance_output, key_terms_output, metadata_output,
                analysis_state, doc_text_state
            ]
        )

        def submit_question_and_clear(doc_text, question, chat_history, lang, analysis_data):
            new_history = ask_enhanced_question(GENAI_CLIENT, GENAI_MODEL_NAME, doc_text, question, chat_history, lang, analysis_data)
            return new_history, ""

        ask_btn.click(
            fn=submit_question_and_clear,
            inputs=[doc_text_state, question_input, chatbot, lang_select, analysis_state],
            outputs=[chatbot, question_input]
        )
        question_input.submit(
            fn=submit_question_and_clear,
            inputs=[doc_text_state, question_input, chatbot, lang_select, analysis_state],
            outputs=[chatbot, question_input]
        )

    return interface

if __name__ == "__main__":
    print("🚀 Starting Advanced AI Legal Document Analyzer...")
    if not PDF2IMAGE_AVAILABLE or not OCR_AVAILABLE:
        print("⚠️ NOTE: Scanned PDFs/images OCR is limited (install pdf2image + pytesseract).")
    app = create_enhanced_interface()
    app.queue().launch(server_name="0.0.0.0", show_error=True)
